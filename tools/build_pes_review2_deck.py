# -*- coding: utf-8 -*-
"""Build the PES-template Phase-3 Review-2 deck (docs/PES_REVIEW2_DECK.pptx).

Follows the department's "Phase-3 Review 2" template slide-for-slide:
  Title · Outline · Suggestions from Review 1 · Proposed Approach / Methodology ·
  Algorithm & Pseudocode · Design Approach · Demonstration and Testing ·
  Applicability of Proposed Methodology in Other Domain · Timeline – Update on
  Pending Tasks · Conclusion · References · Thank You
("Add as many slides as required" — methodology, algorithms, design, demo and
applicability span several slides each.)

Template brand: red 24 pt Trebuchet headings, orange body, project title top-left,
PES logo top-right, rust footer band with the team's names.

Slide text for the content sections lives in docs/review2_content.json (drafted and
adversarially judged against the code and the committed data). Survey / ablation /
fairness numbers are written there as {placeholders} and filled HERE from the same
committed JSONs the papers use, so the deck cannot drift from the data.

Review 2: Tue 13 Oct 2026, 3:00–3:30 pm, B204 — every member presents and is
evaluated individually. Each slide's notes start with its presenter.

Re-run: python tools/build_pes_review2_deck.py
"""
import json
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')
FIG = os.path.join(DOCS, 'figures')
OUT = os.path.join(DOCS, 'PES_REVIEW2_DECK.pptx')
CONTENT = os.path.join(DOCS, 'review2_content.json')


def _load(p):
    with open(p, encoding='utf-8') as f:
        return json.load(f)


SV = _load(os.path.join(DOCS, 'survey_validation.json'))
SX = _load(os.path.join(DOCS, 'survey_extras.json'))
AB = _load(os.path.join(DOCS, 'ablation_results.json'))
AF = _load(os.path.join(DOCS, 'asr_fairness.json'))
inc = SX['incremental']
comp = SX['features']['composites']
chat_ab = {r['config']: r for r in AB['chat']}
full = chat_ab['full recipe (deployed)']
no_conda = chat_ab['- CONDA corpus (domain data)']
fam = AF['by_language_family']


def _paper_pages(name, fallback):
    log = os.path.join(DOCS, name + '.log')
    if os.path.exists(log):
        with open(log, encoding='utf-8', errors='ignore') as f:
            m = re.search(r'Output written on ' + re.escape(name) + r'\.pdf \((\d+) pages', f.read())
        if m:
            return int(m.group(1))
    return fallback


def ci(v, signed=False):
    fmt = '{:+.3f}' if signed else '{:.3f}'
    return '[' + fmt.format(v[0]) + ', ' + fmt.format(v[1]) + ']'


# Every number a content slide may quote as a {placeholder}, bound to the data.
NUM = {
    'rho': f"{SV['construct_validity']['rho']:.3f}",
    'rho_ci': ci(SV['construct_validity']['ci95']),
    'n_scored': str(SV['construct_validity']['n']),
    'n_usable': str(SV['n_usable']),
    'n_raw': str(SV['n_raw']),
    'rho_hours': f"{inc['rho_hours']:.3f}",
    'partial': f"{inc['partial_rho']:.3f}",
    'partial_ci': ci(inc['partial_ci']),
    'pattern': f"{comp['pattern']['rho']:.3f}",
    'volume': f"{comp['volume']['rho']:.3f}",
    'contrast': f"{comp['pattern_minus_volume']['diff']:+.3f}",
    'contrast_ci': ci(comp['pattern_minus_volume']['ci'], signed=True),
    'genre_p': f"{SX['genre']['p']:.3f}",
    'chat_prauc': f"{full['pr_auc']:.3f}",
    'no_conda': f"{no_conda['pr_auc']:.3f}",
    'chat_premise': f"{SX['chat_premise']['rho']:.3f}",
    'clips': f"{AF['overall']['clips']:,}",
    'speech_hours': str(AF['overall']['speech_hours']),
    'wer_drav': f"{fam['Dravidian']['wer'] * 100:.1f}%",
    'wer_tb': f"{fam['Tibeto-Burman']['wer'] * 100:.1f}%",
    'report_pp': str(_paper_pages('PROJECT_PAPER', 48)),
    'ieee_pp': str(_paper_pages('IEEE_PAPER', 6)),
}


def fill(s):
    """Replace {placeholder}s from NUM; any unknown or stray brace is an error, not a silent pass."""
    if not isinstance(s, str):
        return s
    out = re.sub(r'\{(\w+)\}', lambda m: NUM[m.group(1)], s)
    if '{' in out or '}' in out:
        raise ValueError(f'unresolved brace in slide text: {out[:80]!r}')
    return out


# ---------- brand (from the Review 2 template) ----------
RED = RGBColor(0xFF, 0x00, 0x00)        # template headings
ORANGE = RGBColor(0xE0, 0x7B, 0x1A)     # template body / accents
RUST = RGBColor(0xB9, 0x54, 0x27)       # footer band
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x88, 0x88, 0x88)
DARK = RGBColor(0x5A, 0x2E, 0x0E)
LIGHT = RGBColor(0xFB, 0xF3, 0xEA)
CODEBG = RGBColor(0xF7, 0xF4, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Trebuchet MS'
MONO = 'Consolas'
TITLE = 'AI-Driven Gaming Addiction Screening System'
FOOT = 'Kaustubh_Khushee_Kanak_Vidisha'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
LOGO = os.path.join(FIG, 'pes_logo.png')


def _tb(slide, x, y, w, h, text, size=18, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
        font=FONT, italic=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name, r.font.italic = Pt(size), bold, font, italic
        r.font.color.rgb = color
    return box


def _md(s):
    out, bold, buf, i = [], False, '', 0
    while i < len(s):
        if s.startswith('**', i):
            if buf:
                out.append((buf, bold))
            buf, bold, i = '', not bold, i + 2
        else:
            buf += s[i]
            i += 1
    if buf:
        out.append((buf, bold))
    return out


def _rich(p, parts, size, color, font=FONT):
    for t, b in parts:
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold, r.font.name = Pt(size), b, font
        r.font.color.rgb = color


def _band(s):
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.42), W, Inches(0.42))
    band.fill.solid(); band.fill.fore_color.rgb = RUST; band.line.fill.background()
    thin = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.48), W, Inches(0.06))
    thin.fill.solid(); thin.fill.fore_color.rgb = ORANGE; thin.line.fill.background()


def _logo(s):
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, Inches(12.25), Inches(0.12), height=Inches(0.95))


def slide(heading, subheading=None, notes='', presenter=None):
    """Template chrome: project title top-left, logo top-right, red heading, footer band."""
    s = prs.slides.add_slide(BLANK)
    _tb(s, Inches(0.55), Inches(0.18), Inches(8), Inches(0.35), TITLE, 12, False, GREY, font='Arial')
    _logo(s)
    _tb(s, Inches(0.55), Inches(0.72), Inches(11.5), Inches(0.55), heading, 24, False, RED, PP_ALIGN.RIGHT)
    if subheading:
        _tb(s, Inches(0.55), Inches(1.22), Inches(11.5), Inches(0.4), subheading, 15, True, ORANGE, PP_ALIGN.RIGHT)
    _band(s)
    _tb(s, Inches(0), H - Inches(0.36), W, Inches(0.3), FOOT, 10, False, WHITE, PP_ALIGN.CENTER)
    tag = f'[Presenter: {presenter}] ' if presenter else ''
    if notes or tag:
        s.notes_slide.notes_text_frame.text = tag + fill(notes or '')
    return s


def content_top(sd):
    return Inches(1.78) if sd.get('subheading') else Inches(1.5)


# ---------- renderers ----------
def r_bullets(s, sd, size=None):
    if size is None:   # auto-size so six full bullets never reach the footer band
        total = len(fill(sd.get('intro', '')).split()) + sum(len(fill(b).split()) for b in sd.get('bullets', []))
        size = 17 if total <= 110 else (15 if total <= 175 else 14)
    y = content_top(sd)
    box = s.shapes.add_textbox(Inches(0.8), y, Inches(11.8), Inches(6.85) - y)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    if sd.get('intro'):
        p = tf.paragraphs[0]
        first = False
        p.space_after = Pt(10)
        _rich(p, _md(fill(sd['intro'])), size, BLACK)
    for it in sd.get('bullets', []):
        sub = it.startswith('  ')
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(5 if sub else 9)
        p.level = 1 if sub else 0
        _rich(p, _md(('– ' if sub else '▪  ') + fill(it.strip())), size - (3 if sub else 0), BLACK if sub else ORANGE)
    if sd.get('footer'):
        _tb(s, Inches(0.8), Inches(6.45), Inches(11.8), Inches(0.4), fill(sd['footer']), 13, False, DARK, italic=True)


def table(s, rows, x, y, w, col_w=None, size=11, hi=(), grey_body=False):
    nr, nc = len(rows), len(rows[0])
    shp = s.shapes.add_table(nr, nc, x, y, w, Inches(0.38) * nr)
    t = shp.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            tf = cell.text_frame
            tf.word_wrap = True
            col = WHITE if r == 0 else (GREY if grey_body else BLACK)
            _rich(tf.paragraphs[0], _md(fill(str(v))), size, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ORANGE if r == 0 else (LIGHT if (r in hi or r % 2) else WHITE)
    return shp


def r_table(s, sd):
    tb = sd['table']
    rows = [tb['header']] + tb['rows']
    words = sum(len(fill(c).split()) for row in tb['rows'] for c in row)
    size = 12 if words <= 150 else (11 if words <= 240 else 10)
    y = content_top(sd)
    if sd.get('intro'):
        _tb(s, Inches(0.6), y, Inches(12.1), Inches(0.45), fill(sd['intro']), 14, False, BLACK)
        y += Inches(0.5)
    widths = tb.get('col_widths_in') or [12.1 / len(tb['header'])] * len(tb['header'])
    table(s, rows, Inches(0.6), y, Inches(12.1), col_w=[Inches(w) for w in widths], size=size)
    if sd.get('footer'):
        _tb(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.4), fill(sd['footer']), 12, False, DARK, italic=True)


def r_pseudocode(s, sd):
    y = content_top(sd)
    hgt = Inches(6.88) - y
    lines = sd['code']['lines']
    csize = 13 if len(lines) <= 16 else (12 if len(lines) <= 20 else 11)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y, Inches(7.45), hgt)
    panel.fill.solid(); panel.fill.fore_color.rgb = CODEBG
    panel.line.color.rgb = ORANGE; panel.line.width = Pt(1.5)
    box = s.shapes.add_textbox(Inches(0.75), y + Inches(0.08), Inches(7.2), hgt - Inches(0.12))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.space_after = Pt(5)
    _rich(p, [(fill(sd['code']['title']), True)], 13, ORANGE)
    for line in lines:
        p = tf.add_paragraph()
        p.space_after = Pt(0)
        txt = fill(line)
        is_comment = txt.lstrip().startswith('#')
        r = p.add_run()
        r.text = txt
        r.font.size, r.font.name = Pt(csize), MONO
        r.font.color.rgb = GREY if is_comment else BLACK
        r.font.italic = is_comment
    # right panel: plain-words explanation + where it lives in the code
    x = Inches(8.3)
    _tb(s, x, y, Inches(4.4), Inches(0.4), fill(sd.get('side_title') or 'In plain words'), 15, True, ORANGE)
    box = s.shapes.add_textbox(x, y + Inches(0.45), Inches(4.4), hgt - Inches(1.3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(sd.get('side_bullets', [])):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        _rich(p, _md('▪  ' + fill(it)), 12.5, BLACK)
    if sd.get('code_ref'):
        _tb(s, x, Inches(6.28), Inches(4.4), Inches(0.6), 'Code: ' + fill(sd['code_ref']), 10, False, GREY, italic=True)


def r_steps(s, sd):
    steps = sd['steps']
    n = len(steps)
    y = content_top(sd) + Inches(0.05)
    gap = Inches(0.28)
    bw = (Inches(12.1) - gap * (n - 1)) / n
    bh = Inches(3.95)
    for i, st in enumerate(steps):
        x = Inches(0.6) + (bw + gap) * i
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh)
        b.adjustments[0] = 0.06
        b.fill.solid(); b.fill.fore_color.rgb = LIGHT
        b.line.color.rgb = ORANGE; b.line.width = Pt(1.75)
        _tb(s, x + Inches(0.12), y + Inches(0.12), bw - Inches(0.24), Inches(0.35), f'{i + 1}', 20, True, ORANGE)
        _tb(s, x + Inches(0.12), y + Inches(0.55), bw - Inches(0.24), Inches(0.7), fill(st['title']), 14, True, DARK)
        _tb(s, x + Inches(0.12), y + Inches(1.25), bw - Inches(0.24), Inches(1.6), fill(st['body']), 11.5, False, BLACK)
        res = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.08), y + bh - Inches(1.12), bw - Inches(0.16), Inches(1.02))
        res.fill.solid(); res.fill.fore_color.rgb = WHITE
        res.line.color.rgb = RGBColor(0xEE, 0xC9, 0x9E); res.line.width = Pt(0.75)
        _tb(s, x + Inches(0.15), y + bh - Inches(1.07), bw - Inches(0.3), Inches(0.95),
            'Result: ' + fill(st['result']), 10.5, True, RUST)
        if i < n - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.03), y + bh / 2 - Inches(0.14),
                                   gap - Inches(0.06), Inches(0.28))
            a.fill.solid(); a.fill.fore_color.rgb = ORANGE; a.line.fill.background()
    if sd.get('footer'):
        _tb(s, Inches(0.6), y + bh + Inches(0.18), Inches(12.1), Inches(0.6), fill(sd['footer']), 14, False, DARK,
            PP_ALIGN.CENTER, italic=True)


def r_twocol(s, sd):
    y = content_top(sd) + Inches(0.05)
    gap = Inches(0.3)
    cw = (Inches(12.1) - gap) / 2
    ch = Inches(4.45) if sd.get('footer') else Inches(4.95)
    for i, col in enumerate(sd['cols'][:2]):
        x = Inches(0.6) + (cw + gap) * i
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, ch)
        b.adjustments[0] = 0.04
        b.fill.solid(); b.fill.fore_color.rgb = LIGHT
        b.line.color.rgb = ORANGE; b.line.width = Pt(1.75)
        _tb(s, x + Inches(0.2), y + Inches(0.12), cw - Inches(0.4), Inches(0.45), fill(col['title']), 16, True, ORANGE)
        box = s.shapes.add_textbox(x + Inches(0.2), y + Inches(0.62), cw - Inches(0.4), ch - Inches(0.75))
        tf = box.text_frame
        tf.word_wrap = True
        for j, it in enumerate(col.get('bullets', [])):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(7)
            _rich(p, _md('▪  ' + fill(it)), 13, BLACK)
    if sd.get('footer'):
        _tb(s, Inches(0.6), y + ch + Inches(0.12), Inches(12.1), Inches(0.5), fill(sd['footer']), 13, False, DARK,
            PP_ALIGN.CENTER, italic=True)


RENDER = {'bullets': r_bullets, 'table': r_table, 'pseudocode': r_pseudocode, 'steps': r_steps, 'twocol': r_twocol}


def render(sd, presenter):
    s = slide(sd['heading'], sd.get('subheading'), sd.get('notes', ''), presenter)
    RENDER[sd['kind']](s, sd)
    return s


C = _load(CONTENT)        # {track: [slide, ...]} — judged content


def track(key, presenter):
    for sd in C[key]:
        render(sd, presenter)


# =============== 1. Title ===============
s = prs.slides.add_slide(BLANK)
_logo(s)
_tb(s, Inches(0), Inches(1.1), W, Inches(0.7), 'UE23CS441A – Capstone Project Phase – 3', 30, False, BLACK, PP_ALIGN.CENTER)
_tb(s, Inches(0), Inches(1.95), W, Inches(0.6), 'Project Progress Review # 2', 26, False, ORANGE, PP_ALIGN.CENTER)
_tb(s, Inches(2.1), Inches(4.6), Inches(3.0), Inches(2.2),
    'Project Title   :\nProject ID        :\nProject Guide  :\nProject Team   :', 22, False, ORANGE)
# one value box per label line, so the team names never wrap out of alignment
for k, val in enumerate(['AI-Driven Gaming Addiction Screening System', 'PW26_SAS-03', 'Prof. Shridevi Sawant']):
    _tb(s, Inches(5.0), Inches(4.66) + Inches(0.4) * k, Inches(8.0), Inches(0.4), val, 18, False, BLACK)
_tb(s, Inches(5.0), Inches(5.86), Inches(8.0), Inches(0.75),
    'Kaustubh Agarwal (PES1UG23CS291) · Khushee P Kiran (PES1UG23CS303)\n'
    'Kanak Goyal (PES1UG23CS279) · Vidisha Murali (PES1UG23CS681)', 15, False, BLACK)
_band(s)
s.notes_slide.notes_text_frame.text = (
    '[Presenter: Vidisha] Review 2 — Tue 13 Oct 2026, 3:00–3:30 pm, B204, with Prof. Sawant. Everyone presents and is '
    'evaluated individually. One sentence, identical from all four of us: a deployed, multimodal screening system for '
    'parents that measures how a child plays, not just how long — externally validated against a clinical '
    'questionnaire, with its limitations named before you ask.')

# =============== 2. Outline ===============
s = slide('Outline', notes='Read the list in one breath; the order follows the department template.', presenter='Vidisha')
r_bullets(s, {'bullets': [
    'Suggestions from Review 1',
    'Proposed Approach / Methodology',
    'Algorithm & Pseudocode',
    'Design Approach',
    'Demonstration and Testing of the completed modules',
    'Applicability of the proposed methodology in other domains',
    'Timeline – update on pending tasks',
    'Conclusion',
    'References',
]}, size=21)

# =============== 3. Suggestions from Review 1 (team fills in the panel's actual remarks) ===============
s = slide('Suggestions from Review 1',
          notes=('FILL IN BEFORE 13 OCT: one row per remark the Review 1 panel actually made — the remark in their words, '
                 'whether it is feasible (yes / partly / no, and why), and the progress shown since (point to the slide '
                 'number in this deck). Do not leave the grey placeholder text on the slide.'),
          presenter='Vidisha')
_tb(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.45),
    'Remarks from the Review 1 panel, how feasible each is, and the progress we can show today.', 15, False, BLACK)
table(s, [
    ['#', 'Suggestion / remark from the panel', 'Feasibility', 'Progress since Review 1 (where shown in this deck)'],
    ['1', '[panel remark]', '[yes / partly / no — why]', '[what changed · slide no.]'],
    ['2', '[panel remark]', '[yes / partly / no — why]', '[what changed · slide no.]'],
    ['3', '[panel remark]', '[yes / partly / no — why]', '[what changed · slide no.]'],
    ['4', '[panel remark]', '[yes / partly / no — why]', '[what changed · slide no.]'],
], Inches(0.6), Inches(2.05), Inches(12.1), col_w=[Inches(0.5), Inches(5.0), Inches(2.6), Inches(4.0)], size=14,
    grey_body=True)

# =============== 4. Proposed Approach / Methodology ===============
track('methodology', 'Kaustubh')

# =============== 5. Algorithm & Pseudocode ===============
track('algo1_workflow', 'Khushee')
track('algo2_behaviour', 'Kaustubh')
track('algo3_chat', 'Kaustubh')
track('algo4_voice', 'Khushee')
track('algo5_fusion', 'Kanak')

# =============== 6. Design Approach ===============
track('design', 'Kanak')

# =============== 7. Demonstration and Testing ===============
track('demo_status', 'Khushee')
s = slide('Demonstration and Testing of the Modules Completed', 'Results and testing of the completed modules',
          notes=(
              'Model-number follow-ups go to Kaustubh by name. Numbers behind this slide: behaviour macro-F1 0.918, '
              'cross-validation 0.921 ± 0.002, calibration error (ECE) 0.062 → 0.015, screen-time-only baseline 0.702, '
              'five pattern features alone 0.902 vs five volume features 0.871. Chat PR-AUC interval [0.807, 0.841]; '
              'without the in-game chat corpus {no_conda}; 933 Hindi held-out rows never trained on. Voice: chance is '
              '0.25; wav2vec2 headroom 0.776. Survey: ρ {rho} {rho_ci} on {n_scored} people; gain over screen time '
              '+0.167 with interval touching zero (suggestive); partial with hours removed {partial} {partial_ci}; '
              'pattern-minus-volume {contrast} {contrast_ci}; genre test p {genre_p}; one respondent in the '
              'disordered range, so no sensitivity/specificity. Fairness: {speech_hours} h, 117 speakers. Tests: '
              '181 backend (SQLite and Postgres) + 110 Android + 7 paper-vs-data guards, all in CI; 288 concurrent '
              'requests, 0 errors, p50 66 ms.'),
          presenter='Kanak')
table(s, [
    ['Module', 'Result (held-out, real data unless stated)', 'Testing'],
    ['Behaviour model', '**91.6% accuracy** on held-out data. Caveat: the labels are **synthetic** — they came from a simulation grounded on two real surveys, not from clinicians. Calibration brought the score much closer to a true probability.', 'Ablations (retraining with parts removed) with confidence intervals. Screen time alone scores far lower. Pattern (how-you-play) features edge out screen-time features.'],
    ['Chat model', 'On real in-game chat, PR-AUC is **{chat_prauc}**. At our alert threshold (0.95), 95.6% of flagged messages are truly toxic and we catch 42.8% of toxic messages. Hindi precision: 0.968 in Devanagari, 0.958 romanised.', 'Dropping the in-game chat corpus (CONDA) from training cuts PR-AUC sharply. Pretrained toxic-BERT reaches only 0.709, below our simpler model. Hindi test rows were never trained on.'],
    ['Voice model', '**0.574 accuracy** when no speaker appears in both training and test (a speaker-independent split), well above guessing. A random split gave 0.657 — an 8.3-point gap, because the model had memorised voices. We report the honest number.', 'Each speaker stays on one side of every split. A larger model (wav2vec2) scores higher, so headroom exists. Audio augmentation made no difference.'],
    ['**External validation**', 'Our score tracks IGDS9-SF (standard clinical questionnaire): rank correlation ρ **{rho}**, a **moderate** link. Screen-time hours alone: {rho_hours}. How a child plays beats how long: pattern features correlate **{contrast}** higher than screen-time features, a gap distinguishable from zero.', 'Exclusions fixed in advance. Confidence intervals from resampling the same people. Late responses lowered the headline but stay in, by a rule set beforehand. Game genre: no effect. Sensitivity/specificity not computed: only one respondent in the disordered range.'],
    ['Fairness (speech-to-text → toxicity)', '**0 false alerts** across {clips} clips of Indian-English speech, in every accent group. But speech-to-text error (WER, word error rate) varies by first-language family: {wer_drav} for Dravidian speakers up to {wer_tb} for Tibeto-Burman speakers.', 'Svarah, an Indian-English speech corpus, run through the deployed speech-to-text and toxicity scorer.'],
    ['System / apps', 'Backend live on Render and Neon. Signed APKs tested on a phone. Consented family pilot completed. Default path measured on-device: **14% CPU / 288 MB**. Two speech-to-text engines together (dual-STT) **fail our resource gate: default OFF**.', '181 backend tests (SQLite and Postgres), 110 Android tests and 7 paper-versus-data guards run in CI. 288-request load test: 0 errors. Plus fuzzing, CVE and MobSF scans.'],
], Inches(0.6), Inches(1.78), Inches(12.1), col_w=[Inches(1.6), Inches(6.8), Inches(3.7)], size=10, hi=(4,))

# =============== 8. Applicability in other domains ===============
track('applicability', 'Vidisha')

# =============== 9. Timeline – update on pending tasks ===============
s = slide('Timeline – Update on Pending Tasks', 'All modules are complete; these tasks remain',
          notes=('All modules were built, validated and released by September (v2.4.0). Pending: the genre-multiplier '
                 'removal (the one change of approach, on the methodology slide), polishing and submitting the IEEE '
                 'paper, and the department milestones. FINAL REPORT and FINAL EVALUATION dates: fill in when the '
                 'department announces them.'),
          presenter='Kanak')
weeks = ['5 Oct', '12 Oct', '19 Oct', '26 Oct', '2 Nov', '9 Nov', '16 Nov', '23 Nov', '30 Nov']
# (task, owner, first week index, last week index, status) — status: done / review / plan / tba
tasks = [
    ('All modules built, validated, released (Feb – Sep): v2.4.0', 'All', None, None, 'done'),
    ('Review 2 — presentation and demo (13 Oct, 3:00 pm, B204)', 'All', 1, 1, 'review'),
    ('Act on Review 2 feedback', 'All', 1, 3, 'plan'),
    ('Remove the genre multiplier from the served score; version the score', 'Kaustubh, Kanak', 2, 3, 'plan'),
    ('Re-check served score = validated score; update report, papers, decks', 'Kaustubh', 3, 4, 'plan'),
    ('IEEE paper polish: multiple-comparison correction, survey reporting', 'Kaustubh, all', 3, 5, 'plan'),
    ('IEEE conference submission (venue and deadline to be fixed)', 'Kaustubh', 5, 6, 'plan'),
    ('Final project report — submission date: [fill in]', 'All', 2, 7, 'tba'),
    ('Final evaluation / viva — date: [fill in]', 'All', None, None, 'tba'),
]
gx, gy, gw = Inches(0.6), Inches(1.75), Inches(12.1)
label_w, owner_w = Inches(5.1), Inches(1.35)
cell_w = (gw - label_w - owner_w) / len(weeks)
row_h = Inches(0.47)
_tb(s, gx, gy, label_w, row_h, 'Task', 12, True, ORANGE)
_tb(s, gx + label_w, gy, owner_w, row_h, 'Owner', 12, True, ORANGE)
for i, wk in enumerate(weeks):
    _tb(s, gx + label_w + owner_w + cell_w * i, gy, cell_w, row_h, wk, 10, True, ORANGE, PP_ALIGN.CENTER)
for r, (name, owner, a, b, st) in enumerate(tasks, start=1):
    y = gy + row_h * r
    _tb(s, gx, y + Inches(0.06), label_w - Inches(0.1), row_h, name, 10.5, st == 'review', BLACK)
    _tb(s, gx + label_w, y + Inches(0.06), owner_w, row_h, owner, 10.5, False, GREY)
    for i in range(len(weeks)):
        cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, gx + label_w + owner_w + cell_w * i, y, cell_w, row_h - Inches(0.04))
        cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
        cell.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6); cell.line.width = Pt(0.5)
    if st == 'done':
        _tb(s, gx + label_w + owner_w + Inches(0.1), y + Inches(0.08), cell_w * len(weeks), row_h,
            '✓ complete before this timeline', 10.5, True, ORANGE)
    elif a is None:
        _tb(s, gx + label_w + owner_w + Inches(0.1), y + Inches(0.08), cell_w * len(weeks), row_h,
            'date to be announced by the department', 10.5, False, GREY, italic=True)
    else:
        x0 = gx + label_w + owner_w + cell_w * a + Inches(0.05)
        bw = cell_w * (b - a + 1) - Inches(0.1)
        shape = MSO_SHAPE.DIAMOND if st == 'review' else MSO_SHAPE.ROUNDED_RECTANGLE
        if st == 'review':
            x0, bw = gx + label_w + owner_w + cell_w * a + cell_w / 2 - Inches(0.16), Inches(0.32)
        bar = s.shapes.add_shape(shape, x0, y + Inches(0.08), bw, row_h - Inches(0.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = {'review': RED, 'plan': ORANGE, 'tba': RGBColor(0xF2, 0xC2, 0x8A)}[st]
        bar.line.fill.background()
_tb(s, gx, gy + row_h * (len(tasks) + 1) + Inches(0.08), gw, Inches(0.4),
    'Red diamond = Review 2 · Orange = planned team task · Light = department deadline (date to be confirmed) · '
    'weeks start on the date shown.', 10.5, False, GREY)

# =============== 10. Conclusion ===============
track('conclusion', 'Vidisha')

# =============== 11. References (IEEE format; two slides — project sources, then applicability sources) ===============
def refs_slide(title, refs, start, note):
    s = slide(title, notes=note, presenter='Kanak')
    box = s.shapes.add_textbox(Inches(0.6), Inches(1.45), Inches(12.1), Inches(5.45))
    tf = box.text_frame
    tf.word_wrap = True
    for i, ref in enumerate(refs, start):
        p = tf.paragraphs[0] if i == start else tf.add_paragraph()
        p.space_after = Pt(4)
        _rich(p, [(f'[{i}] {fill(ref)}', False)], 11, BLACK)


CORE, DOMAIN = C.get('references') or [], C.get('references_domain') or []
refs_slide('References', CORE, 1, 'IEEE format. Sources the system and its evaluation are built on; the full lists are in the report and the IEEE paper.')
if DOMAIN:
    refs_slide('References (continued)', DOMAIN, len(CORE) + 1,
               'IEEE format. Sources named on the applicability slides, each verified against the publisher or official page.')

# =============== 12. Thank you ===============
s = prs.slides.add_slide(BLANK)
_tb(s, Inches(0), Inches(3.0), W, Inches(1.4), 'Thank You', 40, False, ORANGE, PP_ALIGN.CENTER)
_band(s)

prs.save(OUT)
print(f'wrote {OUT}  ({len(prs.slides)} slides)')
